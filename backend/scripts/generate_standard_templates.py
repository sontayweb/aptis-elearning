import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def generate_docx_template():
    doc = docx.Document()

    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    title = doc.add_heading('BỘ KHUNG MẪU CHUẨN NHẬP ĐỀ THI SPEAKING (PART 2 & PART 3)', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_desc = doc.add_paragraph('Tài liệu mẫu chuẩn hóa dành cho Giáo viên / Ban chuyên môn biên soạn đề Aptis ESOL & General.\nHệ thống sẽ tự động quét bóc tách tranh, từ vựng và câu hỏi theo đúng các thẻ quy ước bên dưới.')
    p_desc.runs[0].font.size = Pt(11)
    p_desc.runs[0].font.italic = True

    doc.add_heading('I. QUY ƯỚC CÁC THẺ ĐỊNH DANH (BẮT BUỘC)', level=1)
    rules = [
        ('=== [ĐỀ MỚI] ===', 'Ký hiệu bắt đầu một đề thi mới. Đặt ở đầu mỗi đề.'),
        ('[PHẦN THI]:', 'Điền "Speaking Part 2" hoặc "Speaking Part 3".'),
        ('[CHỦ ĐỀ]:', 'Tên chủ đề thực tế của bức tranh (VD: Làm việc văn phòng, Cắm trại dã ngoại...).'),
        ('[HÌNH ẢNH]:', 'Chèn trực tiếp 1 ảnh (đối với Part 2) hoặc 2 ảnh (đối với Part 3) ngay dưới thẻ này.'),
        ('[TỪ VỰNG]:', 'Chèn bảng 2 cột gồm English và Tiếng Việt bên dưới thẻ này.'),
        ('[CÂU HỎI 1]:', 'Nội dung câu hỏi số 1.'),
        ('[CÂU HỎI 2]:', 'Nội dung câu hỏi số 2.'),
        ('[CÂU HỎI 3]:', 'Nội dung câu hỏi số 3.')
    ]
    for tag, desc in rules:
        p = doc.add_paragraph()
        r1 = p.add_run(tag + ' ')
        r1.bold = True
        r1.font.color.rgb = RGBColor(30, 64, 175)
        p.add_run(desc)

    # Đề thi 1
    doc.add_heading('II. ĐỀ THI MẪU 1: SPEAKING PART 2 (1 TRANH MÔ TẢ)', level=1)
    doc.add_paragraph('=== [ĐỀ MỚI] ===').runs[0].bold = True
    doc.add_paragraph('[PHẦN THI]: Speaking Part 2')
    doc.add_paragraph('[CHỦ ĐỀ]: Nhóm người trong phòng họp công ty')
    doc.add_paragraph('[HÌNH ẢNH]: (Chèn 1 bức ảnh chụp tại đây)')

    doc.add_paragraph('[TỪ VỰNG]:')
    table2 = doc.add_table(rows=1, cols=2)
    table2.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr2 = table2.rows[0].cells
    hdr2[0].text = 'English'
    hdr2[1].text = 'Tiếng Việt'
    hdr2[0].paragraphs[0].runs[0].bold = True
    hdr2[1].paragraphs[0].runs[0].bold = True

    v2 = [
        ('office', 'văn phòng'),
        ('colleagues', 'đồng nghiệp'),
        ('meeting room', 'phòng họp'),
        ('discussing', 'thảo luận'),
        ('laptop screen', 'màn hình máy tính')
    ]
    for en, vi in v2:
        r = table2.add_row().cells
        r[0].text = en
        r[1].text = vi

    doc.add_paragraph()
    doc.add_paragraph('[CÂU HỎI 1]: Describe what you see in the picture.')
    doc.add_paragraph('[CÂU HỎI 2]: Tell me about a time when you had to work in a team.')
    doc.add_paragraph('[CÂU HỎI 3]: Why is team cooperation important in the modern workplace?')

    # Đề thi 2
    doc.add_heading('III. ĐỀ THI MẪU 2: SPEAKING PART 3 (CẶP TRANH ĐỐI CHIẾU SO SÁNH)', level=1)
    doc.add_paragraph('=== [ĐỀ MỚI] ===').runs[0].bold = True
    doc.add_paragraph('[PHẦN THI]: Speaking Part 3')
    doc.add_paragraph('[CHỦ ĐỀ]: So sánh kỳ nghỉ leo núi mạo hiểm và nghỉ dưỡng bãi biển')
    doc.add_paragraph('[HÌNH ẢNH]: (Chèn 2 bức ảnh đối chiếu tại đây: Tranh A và Tranh B)')

    doc.add_paragraph('[TỪ VỰNG]:')
    table3 = doc.add_table(rows=1, cols=2)
    table3.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr3 = table3.rows[0].cells
    hdr3[0].text = 'English'
    hdr3[1].text = 'Tiếng Việt'
    hdr3[0].paragraphs[0].runs[0].bold = True
    hdr3[1].paragraphs[0].runs[0].bold = True

    v3 = [
        ('mountain climbing / hiking', 'leo núi / đi bộ dã ngoại'),
        ('tropical beach', 'bãi biển nhiệt đới'),
        ('adventure & challenge', 'phiêu lưu và thử thách'),
        ('relaxation & peaceful', 'thư giãn và yên bình'),
        ('physical endurance', 'sức bền thể lực')
    ]
    for en, vi in v3:
        r = table3.add_row().cells
        r[0].text = en
        r[1].text = vi

    doc.add_paragraph()
    doc.add_paragraph('[CÂU HỎI 1]: Tell me what you see in the two pictures and compare them.')
    doc.add_paragraph('[CÂU HỎI 2]: Which of these two types of holidays would you prefer to spend with your family?')
    doc.add_paragraph('[CÂU HỎI 3]: How do modern travel habits affect natural landscapes?')

    out_docx = 'd:/sontayweb/aptis-elearning/frontend/public/templates/Mau_Nhap_De_Speaking_Part2_Part3.docx'
    doc.save(out_docx)
    print(f"✅ Đã tạo thành công Word Template: {out_docx}")

def generate_excel_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reading_Listening_Template"

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    headers = [
        ("Skill", 12),
        ("Part", 8),
        ("Topic", 25),
        ("Passage_Text", 35),
        ("Question_No", 12),
        ("Prompt", 35),
        ("Option_A", 20),
        ("Option_B", 20),
        ("Option_C", 20),
        ("Option_D", 20),
        ("Correct_Answer", 15),
        ("Explanation", 30)
    ]

    for col_idx, (h_name, width) in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=h_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    sample_rows = [
        ("READING", 1, "Thư mời tiệc sinh nhật", "Dear Sarah, I am writing to invite you to my birthday party next Sunday...", 1, "The party will be held at _______ house.", "my", "mine", "me", "I", "my", "Tính từ sở hữu đứng trước danh từ house"),
        ("READING", 1, "Thư mời tiệc sinh nhật", "", 2, "Please let me know if you can _______ by Friday.", "come", "came", "coming", "comes", "come", "Sau can dùng động từ nguyên thể"),
        ("LISTENING", 1, "Thông báo chuyển ga tàu", "Attention all passengers for Train 42B to Oxford. The departure platform has been changed to Platform 5 due to maintenance.", 1, "Which platform is the train departing from?", "Platform 2", "Platform 4", "Platform 5", "Platform 9", "Platform 5", "Thông báo nêu rõ Platform 5"),
        ("LISTENING", 2, "Dự báo thời tiết cuối tuần", "Good evening. Looking at this weekend, Saturday will start with bright sunshine, but heavy rain is expected on Sunday afternoon.", 1, "What will the weather be like on Sunday afternoon?", "Sunny", "Snowy", "Heavy rain", "Foggy", "Heavy rain", "Nêu rõ heavy rain vào Sunday afternoon")
    ]

    for row_idx, row_data in enumerate(sample_rows, 2):
        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.alignment = center_align if col_idx in [1, 2, 5, 11] else left_align
            cell.border = thin_border

    out_xlsx = 'd:/sontayweb/aptis-elearning/frontend/public/templates/Mau_Nhap_De_Reading_Listening.xlsx'
    wb.save(out_xlsx)
    print(f"✅ Đã tạo thành công Excel Template: {out_xlsx}")

if __name__ == '__main__':
    generate_docx_template()
    generate_excel_template()
