import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def generate_writing_docx():
    doc = docx.Document()

    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    title = doc.add_heading('BỘ KHUNG MẪU CHUẨN NHẬP ĐỀ THI WRITING (4 PHẦN CHUẨN APTIS ESOL)', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_desc = doc.add_paragraph('Tài liệu mẫu chuẩn hóa dành cho Giáo viên / Ban chuyên môn biên soạn đề Aptis Writing (50 phút).\nHệ thống sẽ tự động quét bóc tách ngữ cảnh, câu hỏi và bảng từ vựng theo đúng các thẻ quy ước bên dưới.')
    p_desc.runs[0].font.size = Pt(11)
    p_desc.runs[0].font.italic = True

    doc.add_heading('I. QUY ƯỚC CÁC THẺ ĐỊNH DANH KỸ NĂNG VIẾT (WRITING)', level=1)
    rules = [
        ('=== [ĐỀ MỚI] ===', 'Ký hiệu bắt đầu một đề thi mới. Đặt ở đầu mỗi đề.'),
        ('[PHẦN THI]:', 'Điền "Writing Part 1", "Writing Part 2", "Writing Part 3" hoặc "Writing Part 4".'),
        ('[CHỦ ĐỀ]:', 'Tên chủ đề hoặc tên Câu Lạc Bộ (VD: Câu lạc bộ Thể thao, Book Club...).'),
        ('[NGỮ CẢNH]:', 'Nội dung thông báo của câu lạc bộ (Dành riêng cho Part 4 viết thư/email).'),
        ('[TỪ VỰNG]:', 'Chèn bảng 2 cột gồm English và Tiếng Việt bên dưới thẻ này.'),
        ('[CÂU HỎI 1]:', 'Nội dung câu hỏi số 1.'),
        ('[CÂU HỎI 2]:', 'Nội dung câu hỏi số 2.'),
        ('[CÂU HỎI 3]:', 'Nội dung câu hỏi số 3... (tùy theo số lượng câu của từng Part).')
    ]
    for tag, desc in rules:
        p = doc.add_paragraph()
        r1 = p.add_run(tag + ' ')
        r1.bold = True
        r1.font.color.rgb = RGBColor(30, 64, 175)
        p.add_run(desc)

    # Part 1
    doc.add_heading('II. ĐỀ MẪU 1: WRITING PART 1 (ĐIỀN FORM - 5 CÂU HỎI TỪ 1 ĐẾN 5 TỪ)', level=1)
    doc.add_paragraph('=== [ĐỀ MỚI] ===').runs[0].bold = True
    doc.add_paragraph('[PHẦN THI]: Writing Part 1')
    doc.add_paragraph('[CHỦ ĐỀ]: Đơn gia nhập Câu Lạc Bộ Thể Thao (Sports Club)')
    doc.add_paragraph('[TỪ VỰNG]:')
    t1 = doc.add_table(rows=1, cols=2)
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    h1 = t1.rows[0].cells
    h1[0].text = 'English'
    h1[1].text = 'Tiếng Việt'
    h1[0].paragraphs[0].runs[0].bold = True
    h1[1].paragraphs[0].runs[0].bold = True
    for en, vi in [('occupation', 'nghề nghiệp'), ('leisure activity', 'hoạt động giải trí'), ('frequency', 'tần suất')]:
        r = t1.add_row().cells
        r[0].text = en
        r[1].text = vi

    doc.add_paragraph()
    doc.add_paragraph('[CÂU HỎI 1]: What is your primary occupation?')
    doc.add_paragraph('[CÂU HỎI 2]: Where do you currently live?')
    doc.add_paragraph('[CÂU HỎI 3]: What is your favorite sport or physical activity?')
    doc.add_paragraph('[CÂU HỎI 4]: How often do you exercise during the week?')
    doc.add_paragraph('[CÂU HỎI 5]: Why do you want to join our Sports Club?')

    # Part 2
    doc.add_heading('III. ĐỀ MẪU 2: WRITING PART 2 (VIẾT NGẮN THÔNG TIN CÁ NHÂN - 20 ĐẾN 30 TỪ)', level=1)
    doc.add_paragraph('=== [ĐỀ MỚI] ===').runs[0].bold = True
    doc.add_paragraph('[PHẦN THI]: Writing Part 2')
    doc.add_paragraph('[CHỦ ĐỀ]: Chia sẻ lý do gia nhập Sports Club')
    doc.add_paragraph('[TỪ VỰNG]:')
    t2 = doc.add_table(rows=1, cols=2)
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    h2 = t2.rows[0].cells
    h2[0].text = 'English'
    h2[1].text = 'Tiếng Việt'
    h2[0].paragraphs[0].runs[0].bold = True
    h2[1].paragraphs[0].runs[0].bold = True
    for en, vi in [('stay fit', 'giữ dáng khỏe mạnh'), ('expand social circle', 'mở rộng quan hệ bạn bè')]:
        r = t2.add_row().cells
        r[0].text = en
        r[1].text = vi

    doc.add_paragraph()
    doc.add_paragraph('[CÂU HỎI 1]: Please tell us why you are interested in joining our Sports Club and what you hope to achieve here (20-30 words).')

    # Part 3
    doc.add_heading('IV. ĐỀ MẪU 3: WRITING PART 3 (TƯƠNG TÁC PHÒNG CHAT CLB - 3 CÂU, 30-40 TỪ/CÂU)', level=1)
    doc.add_paragraph('=== [ĐỀ MỚI] ===').runs[0].bold = True
    doc.add_paragraph('[PHẦN THI]: Writing Part 3')
    doc.add_paragraph('[CHỦ ĐỀ]: Thảo luận phòng chat thành viên Sports Club')
    doc.add_paragraph('[TỪ VỰNG]:')
    t3 = doc.add_table(rows=1, cols=2)
    t3.alignment = WD_TABLE_ALIGNMENT.CENTER
    h3 = t3.rows[0].cells
    h3[0].text = 'English'
    h3[1].text = 'Tiếng Việt'
    h3[0].paragraphs[0].runs[0].bold = True
    h3[1].paragraphs[0].runs[0].bold = True
    for en, vi in [('teamwork', 'tinh thần đồng đội'), ('individual workout', 'tập luyện cá nhân'), ('outdoor event', 'sự kiện ngoài trời')]:
        r = t3.add_row().cells
        r[0].text = en
        r[1].text = vi

    doc.add_paragraph()
    doc.add_paragraph('[CÂU HỎI 1]: Member Sam: Welcome to our sports group! How long have you been interested in fitness and sports?')
    doc.add_paragraph('[CÂU HỎI 2]: Member Linda: Do you prefer exercising outdoors in the park or going to a modern gym? Why?')
    doc.add_paragraph('[CÂU HỎI 3]: Member Alex: We want to organize a charity running event next month. What suggestions do you have for us?')

    # Part 4
    doc.add_heading('V. ĐỀ MẪU 4: WRITING PART 4 (VIẾT EMAIL BẠN BÈ VÀ EMAIL TRANG TRỌNG)', level=1)
    doc.add_paragraph('=== [ĐỀ MỚI] ===').runs[0].bold = True
    doc.add_paragraph('[PHẦN THI]: Writing Part 4')
    doc.add_paragraph('[CHỦ ĐỀ]: Thông báo đóng cửa trung tâm thể thao bảo trì')
    doc.add_paragraph('[NGỮ CẢNH]: Dear Members, Due to emergency electrical maintenance, our central sports complex will be closed for the next two months. All upcoming weekend tournaments have been postponed until further notice. Sincerely, Sports Club Management.')
    doc.add_paragraph('[TỪ VỰNG]:')
    t4 = doc.add_table(rows=1, cols=2)
    t4.alignment = WD_TABLE_ALIGNMENT.CENTER
    h4 = t4.rows[0].cells
    h4[0].text = 'English'
    h4[1].text = 'Tiếng Việt'
    h4[0].paragraphs[0].runs[0].bold = True
    h4[1].paragraphs[0].runs[0].bold = True
    for en, vi in [('renovation / maintenance', 'bảo trì / sửa chữa'), ('postponed', 'hoãn lại'), ('constructive suggestion', 'đề xuất mang tính xây dựng')]:
        r = t4.add_row().cells
        r[0].text = en
        r[1].text = vi

    doc.add_paragraph()
    doc.add_paragraph('[CÂU HỎI 1]: Write an email to your friend who is also a club member. Express your feelings about this news and suggest an alternative place where both of you can exercise together (about 50 words).')
    doc.add_paragraph('[CÂU HỎI 2]: Write an email to the Club President. Express your disappointment, explain how the closure affects your fitness routine, and propose reasonable solutions for members during the closure (120-150 words).')

    target_dir = 'd:/sontayweb/aptis-elearning/frontend/public/templates'
    os.makedirs(target_dir, exist_ok=True)
    out_docx = os.path.join(target_dir, 'Mau_Nhap_De_Writing.docx')
    doc.save(out_docx)
    print(f"SUCCESS: Generated {out_docx}")

if __name__ == '__main__':
    generate_writing_docx()
